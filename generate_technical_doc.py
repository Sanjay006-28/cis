#!/usr/bin/env python3
"""
Technical Specification & Architecture Documentation Generator
for Android Malware Detection System.
Generates:
  1. PROJECT_TECHNICAL_SPECIFICATION.md
  2. Android_Malware_Detection_Technical_Specification.docx
  3. Android_Malware_Detection_Technical_Specification.pdf
"""

import os
import sys
from pathlib import Path

# Paths
BASE_DIR = Path(__file__).resolve().parent
STATIC_OUT = BASE_DIR / "static_outputs"
DYNAMIC_OUT = BASE_DIR / "dynamic_outputs_v2"

# -----------------------------------------------------------------------------
# 1. MARKDOWN GENERATION
# -----------------------------------------------------------------------------
MD_CONTENT = """# Android Malware Detection System: Architecture, Design Flow & Technical Specifications

**Project Title:** AI-Based Detection of Malicious Sideloaded Android Applications Using Runtime Threat Intelligence and Dynamic Emulation  
**Document Version:** 2.0 (Production & Academic Specification)  
**Target Environment:** Cross-Platform (Windows 11 / WSL 2 Ubuntu / Android Virtual Device)  

---

## Executive Summary

Sideloaded Android applications—installed outside curated repositories like Google Play—bypass pre-deployment vetting, exposing users to sophisticated adware, spyware, banking Trojans, and root exploits. Modern Android malware frequently leverages runtime obfuscation, dynamic payload loading, and anti-analysis evasion tactics that defeat naive static signature scanners.

This project delivers an end-to-end, multi-stage defense-in-depth framework featuring:
1. **Static Analysis Engine (Drebin-215)**: Safe APK bytecode decompilation extracting 215 permissions, intents, API call signatures, and sensitive shell commands.
2. **Machine Learning Classifiers & Ensemble**: Random Forest and XGBoost with soft probability voting achieving **98.14% accuracy** on held-out static test benchmarks.
3. **Local Explainable AI (SHAP)**: TreeExplainer-driven local feature attribution identifying exact structural indicators driving each classification verdict.
4. **Automated Threat Intelligence Enrichment**: Extraction of embedded public IP addresses and domain hostnames cross-referenced in real-time against the AbuseIPDB REST API v2.
5. **Dynamic Sandbox Emulation (CICMalDroid2020 & strace)**: Live capture of Linux kernel system calls on a rooted Android emulator with UI monkey event fuzzing, evaluated against deep learning models (TCN, BiLSTM, Transformer) and XGBoost.
6. **Interactive Streamlit Web Dashboard**: Production-grade operational console providing single-click APK triage, live ADB diagnostic verification, SHAP risk visualization, and downloadable JSON security audit reports.

---

## 1. System Architecture

The system implements a decoupled, five-tier architecture ensuring isolation, fault tolerance, and modular scalability.

```mermaid
flowchart TD
    subgraph UI_Tier["1. Presentation & Operational Tier (Streamlit Web App)"]
        UI[Streamlit Dashboard app.py]
        UI --> Tab1[Tab 1: Static APK Analysis & Threat Intel]
        UI --> Tab2[Tab 2: Static Model Evaluation Benchmarks]
        UI --> Tab3[Tab 3: Dynamic Model Evaluation Benchmarks]
        UI --> Tab4[Tab 4: Live Dynamic ADB Emulation Sandbox]
        UI --> Tab5[Tab 5: System Architecture & Disclosures]
    end

    subgraph Static_Tier["2. Static Bytecode Ingestion & Feature Extraction"]
        APK_In[Uploaded / Sideloaded APK] --> Androguard[Androguard 4.x Bytecode & Manifest Parser]
        Androguard --> ManifestExtract[AndroidManifest.xml: Permissions & Intent Filters]
        Androguard --> DexExtract[classes.dex: Bytecode Strings, API Method Signatures, Types]
        ManifestExtract & DexExtract --> FeatureVector[215-Dimensional Binary Drebin Vector]
        DexExtract --> NetHarvest[Embedded Public IPs & Hostnames]
    end

    subgraph ML_Static_Tier["3. Static Inference & Explainable AI"]
        FeatureVector --> RF_Model[Random Forest Classifier]
        FeatureVector --> XGB_Model[XGBoost Classifier]
        RF_Model & XGB_Model --> SoftEnsemble["Soft Probability Voting: P = 0.5*(P_RF + P_XGB)"]
        XGB_Model --> SHAP_Explainer[SHAP TreeExplainer]
        SHAP_Explainer --> LocalAttribution[Local Feature Attribution & Bar Chart]
    end

    subgraph Intel_Tier["4. Live Threat Intelligence Subsystem"]
        NetHarvest --> IP_Filter[RFC 1918 Private IP Filtering via ipaddress]
        IP_Filter --> DNS_Resolver[Live Socket DNS Resolution socket.getaddrinfo]
        DNS_Resolver --> AbuseIPDB[AbuseIPDB REST API v2 Endpoint]
        AbuseIPDB --> ThreatScore[Abuse Confidence Scores & Risk Thresholding >= 50%]
    end

    subgraph Dynamic_Tier["5. Dynamic Capture & Kernel Syscall Sandbox"]
        LiveAPK[APK File] --> ADB_Bridge[ADB Client-Server Bridge]
        ADB_Bridge --> Emulator["Rooted Android Emulator (API 34 / Google APIs)"]
        Emulator --> MonkeyFuzzer[UI Exerciser Monkey: Touch, Gestures, Key Events]
        Emulator --> KernelStrace["strace -c -f -p PID (Counts 139 Linux Kernel Syscalls)"]
        KernelStrace --> SyscallParser[Regex Table Parser & Syscall Frequency Normalizer]
        SyscallParser --> DynXGB[Dynamic XGBoost & Deep Learning Models TCN/BiLSTM/Transformer]
        DynXGB --> DynVerdict[Dynamic Malware Probability & Syscall SHAP]
    end

    SoftEnsemble --> UI
    LocalAttribution --> UI
    ThreatScore --> UI
    DynVerdict --> UI
```

---

## 2. Design Flow & Lifecycle

The end-to-end execution lifecycle follows a sequential state transition pipeline:

```mermaid
sequenceDiagram
    autonumber
    actor Analyst as Security Analyst
    participant Web as Streamlit Frontend (app.py)
    participant StaticEngine as Static Engine (phase3_live_demo.py)
    participant ThreatAPI as AbuseIPDB API v2
    participant DynEngine as Dynamic Sandbox (phase5_live_capture.py)
    participant Emulator as Android Emulator & ADB

    Analyst->>Web: Uploads APK / Selects Sample APK
    
    rect rgb(240, 248, 255)
        note over Web,StaticEngine: Static Analysis Phase
        Web->>StaticEngine: extract_apk_info(apk_path)
        StaticEngine->>StaticEngine: Parse Manifest + DEX bytecode safely
        StaticEngine->>StaticEngine: Match 215 Drebin features
        StaticEngine->>StaticEngine: Predict via RF + XGBoost Soft Ensemble
        StaticEngine->>StaticEngine: Compute SHAP TreeExplainer attributions
        StaticEngine-->>Web: Static Verdict, Confidence %, Top SHAP Features
    end

    rect rgb(255, 245, 238)
        note over Web,ThreatAPI: Threat Intelligence Phase
        StaticEngine->>StaticEngine: Harvest IP & domain regexes
        StaticEngine->>ThreatAPI: Query reputation for public IPs
        ThreatAPI-->>StaticEngine: Abuse confidence score, country, report count
        StaticEngine-->>Web: Network Threat Report & Indicators
    end

    rect rgb(245, 255, 245)
        note over Web,Emulator: Dynamic Emulation Phase (Optional / Live)
        Analyst->>Web: Click "Run Live Dynamic Analysis"
        Web->>DynEngine: Subprocess invocation with duration (e.g. 30s)
        DynEngine->>Emulator: Check root & push strace binary
        DynEngine->>Emulator: adb install -r -g apk_path
        DynEngine->>Emulator: Launch app intent & poll PID
        DynEngine->>Emulator: Attach strace -c -f -p PID & run Monkey events
        DynEngine->>Emulator: Terminate trace, pull strace_out.txt, uninstall APK
        DynEngine->>DynEngine: Parse table, normalize frequencies, run Dynamic XGBoost
        DynEngine-->>Web: JSON payload (_p5_last_result.json)
    end

    Web->>Analyst: Render Unified Security Dashboard & JSON Download
```

---

## 3. Data Flow Architecture (DFD)

### DFD Level 0: Context Diagram
```mermaid
flowchart LR
    APK_File[Raw Android APK File] --> System((Android Malware Detection System))
    Threat_Feeds[AbuseIPDB Threat Feed] <--> System
    System --> Security_Report[Security Verdict, SHAP Explanations, Risk Metrics, Audit JSON]
    Security_Auditor[Security Analyst] -->|Upload APK & Configure Timeout| System
```

### DFD Level 1: Subsystem Data Flow
```mermaid
flowchart TD
    P1[1.0 Static Extraction] -->|215-D Feature Vector| P2[2.0 Static Ensemble Inference]
    P1 -->|Extracted IPs & Domains| P3[3.0 Threat Intelligence Resolution]
    P2 -->|SHAP Values & Probabilities| P5[5.0 Result Aggregation & Reporting]
    P3 -->|Abuse Scores & Geo-Data| P5
    P4[4.0 Dynamic Sandbox Emulation] -->|139 Syscall Frequency Vector| P6[6.0 Dynamic ML/DL Inference]
    P6 -->|Dynamic Verdict & Syscall SHAP| P5
    P5 --> Output[Rendered Web UI & JSON Security Report]
```

---

## 4. Machine Learning & Deep Learning Algorithms

### 4.1 Static Analysis Pipeline (Drebin-215)
- **Dataset:** Drebin-215 benchmark comprising 15,036 total application feature vectors (5,560 Malware, 9,476 Benign) deduplicated to 7,254 unique samples.
- **Feature Space:** 215 binary indicators across 4 domains:
  1. *Permissions* (e.g., `SEND_SMS`, `READ_PHONE_STATE`, `ACCESS_FINE_LOCATION`)
  2. *Intents/Action Filters* (e.g., `BOOT_COMPLETED`, `PACKAGE_REPLACED`)
  3. *Sensitive API Calls* (e.g., `getDeviceId`, `sendTextMessage`, `DexClassLoader`)
  4. *Shell & Root Commands* (e.g., `chmod`, `mount`, `chown`, `/system/bin`)
- **Classifiers Used:**
  - **Random Forest:** Ensemble of 100 de-correlated decision trees with Gini impurity splitting.
  - **XGBoost:** Extreme Gradient Boosted decision trees with shrinkage (`learning_rate=0.1`), max depth 6, and log-loss objective.
  - **Soft Voting Ensemble:** Fused prediction computing the arithmetic mean of calibrated probabilities:
    $$P_{\text{ensemble}}(\text{Malware} \mid x) = \frac{1}{2} \left[ P_{\text{RF}}(\text{Malware} \mid x) + P_{\text{XGB}}(\text{Malware} \mid x) \right]$$

### 4.2 Dynamic Analysis Pipeline (CICMalDroid2020)
- **Dataset:** CICMalDroid2020 dynamic behavioral dataset with 11,520 applications traced during runtime execution.
- **Feature Space:** 139 Linux kernel system call frequency counts normalized to capture process execution behavior.
- **Model Architectures Benchmarked:**
  1. **XGBoost Classifier:** Tree-boosted model achieving superior tabular discrimination on syscall frequency vectors.
  2. **Multi-Branch Temporal Convolutional Network (TCN):** Dilated causal 1D convolutions with residual skip connections and receptive field expansion for sequence dynamics.
  3. **Bidirectional LSTM (BiLSTM):** Recurrent neural network capturing forward and backward temporal dependencies across execution steps.
  4. **Transformer:** Multi-head self-attention encoder modeling non-local syscall interaction dependencies.

### 4.3 Validation Strategy & Rigor
- **Cluster-Held-Out Cross-Validation:** To prevent artificial performance inflation caused by family-level data leakage (where variants of the same malware family exist in both train and test splits), samples were clustered by behavioral family and entire clusters were held out during evaluation.
- **McNemar Statistical Significance Tests:** Pairwise contingency table tests with Bonferroni correction confirming whether accuracy differences between deep learning architectures and XGBoost are statistically significant ($p < 0.05$).

---

## 5. Threat Intelligence Subsystem

The runtime threat intelligence pipeline bridges static code inspection with active perimeter defense:

1. **Bytecode Mining:** Scans string pools across all compiled `classes.dex` files using high-precision IPv4 and RFC-compliant domain regular expressions:
   - IPv4 regex: `\b(?:\d{1,3}\.){3}\d{1,3}\b`
   - Domain regex: `\b(?:[a-zA-Z0-9-]+\.)+[a-zA-Z]{2,}\b`
2. **RFC 1918 & Known Host Filtering:** Drops non-routable private addresses (`10.0.0.0/8`, `172.16.0.0/12`, `192.168.0.0/16`, `127.0.0.0/8`) and legitimate development namespaces (e.g., `schemas.android.com`, `googleapis.com`, `google.com`).
3. **Optional DNS Resolution:** Resolves embedded hostnames to public IPv4 targets using `socket.getaddrinfo`.
4. **AbuseIPDB Integration:** Queries the AbuseIPDB v2 endpoint (`https://api.abuseipdb.com/api/v2/check`) with 90-day historical window. Captures:
   - Abuse Confidence Score (0% to 100%)
   - Total abuse report counts and reporting country
   - ISP / Usage type classification
5. **Threat Scoring:** Any IP with $\ge 50\%$ confidence is flagged as an active Malicious Command-and-Control (C2) / exfiltration endpoint.

---

## 6. Explainable AI (XAI) Framework

To satisfy regulatory and security auditing standards, every prediction is paired with local feature explanations via **SHAP (SHapley Additive exPlanations)** based on cooperative game theory:

$$\phi_i(x) = \sum_{S \subseteq F \setminus \{i\}} \frac{|S|! (|F| - |S| - 1)!}{|F|!} \left[ f_x(S \cup \{i\}) - f_x(S) \right]$$

- **TreeExplainer:** Evaluates feature contributions in polynomial time $O(TLD^2)$ rather than exponential sampling time.
- **Local Attribution:** Highlights top features increasing malware probability (red/warning) versus features establishing benign profile (blue/green).
- **Global Interpretability:** Summary beeswarm plots and mean absolute SHAP bar charts reveal global feature importance rankings across the entire training corpus.

---

## 7. Experimental Results & Performance Benchmarks

### 7.1 Static Analysis Evaluation (Drebin-215 Test Split)
Evaluated on a 20% held-out test set ($N = 1,451$ unique applications):

| Model Architecture | Accuracy | Precision | Recall (Malware) | F1-Score | ROC-AUC |
| :--- | :---: | :---: | :---: | :---: | :---: |
| **Random Forest** | 97.11% | 96.02% | 91.55% | 0.9373 | 0.9926 |
| **XGBoost** | 98.00% | 95.91% | **95.63%** | 0.9577 | **0.9961** |
| **Soft Voting Ensemble (RF + XGB)** | **98.14%** | **96.75%** | 95.34% | **0.9604** | 0.9952 |

5-Fold Stratified Cross-Validation confirms model stability:
- **Random Forest:** $97.20\% \pm 0.22\%$ accuracy, ROC-AUC $0.9930 \pm 0.0012$
- **XGBoost:** $97.63\% \pm 0.36\%$ accuracy, ROC-AUC $0.9952 \pm 0.0017$
- **Fused Ensemble:** $97.63\% \pm 0.33\%$ accuracy, ROC-AUC $0.9948 \pm 0.0012$

### 7.2 Dynamic System-Call Evaluation (CICMalDroid2020)
Evaluated on 11,520 applications across 139 system calls:

#### Standard Random Split (80/20):
| Model | Test Accuracy | Precision | Recall | Specificity | F1-Score | ROC-AUC |
| :--- | :---: | :---: | :---: | :---: | :---: | :---: |
| **XGBoost** | **97.31%** | **99.01%** | **97.79%** | **94.69%** | **0.9840** | **0.9915** |
| **TCN (Multi-Branch)** | 94.53% | 96.33% | 97.23% | 79.89% | 0.9678 | 0.9554 |
| **Transformer** | 90.93% | 97.96% | 91.16% | 89.66% | 0.9444 | 0.9691 |
| **BiLSTM** | 90.80% | 97.33% | 91.62% | 86.31% | 0.9439 | 0.9523 |

#### Cluster-Held-Out Cross-Validation (Conservative / Zero-Leakage):
| Model | Accuracy | Precision | Recall | Specificity | F1-Score | ROC-AUC |
| :--- | :---: | :---: | :---: | :---: | :---: | :---: |
| **XGBoost** | **91.99% ± 3.36%** | **97.48% ± 1.49%** | **92.98% ± 4.12%** | **87.49% ± 5.13%** | **0.9513** | **0.9600** |
| **Transformer** | 86.15% ± 6.62% | 96.99% ± 1.61% | 86.32% ± 8.64% | 85.60% ± 6.13% | 0.9112 | 0.9419 |
| **TCN (Multi-Branch)** | 84.99% ± 3.14% | 96.33% ± 2.19% | 85.54% ± 3.89% | 82.74% ± 9.29% | 0.9056 | 0.9013 |
| **BiLSTM** | 79.69% ± 9.25% | 95.56% ± 1.72% | 79.74% ± 11.40% | 79.31% ± 8.74% | 0.8657 | 0.8806 |

### 7.3 McNemar Statistical Significance Matrix
Pairwise McNemar hypothesis tests on cluster-held-out predictions confirm significant performance divergence:
- **TCN vs. BiLSTM:** $p = 0.0$ ($1481$ TCN-only correct vs $870$ BiLSTM-only correct) $\rightarrow$ **Significant**
- **TCN vs. Transformer:** $p = 0.0011$ ($758$ TCN-only correct vs $891$ Transformer-only correct) $\rightarrow$ **Significant**
- **TCN vs. XGBoost:** $p = 0.0$ ($587$ TCN-only correct vs $1393$ XGBoost-only correct) $\rightarrow$ **Significant**

---

## 8. Technical Specifications & Environment Stack

| Layer | Component | Version / Specification | Role |
| :--- | :--- | :--- | :--- |
| **Host OS** | Windows 11 64-bit | Build 22621+ | Primary developer workstation & Android Emulator host |
| **Runtime Subsystem** | WSL 2 (Ubuntu 24.04 LTS) | Linux Kernel 5.15+ | Isolated Linux environment running Python & Streamlit |
| **Programming Language** | Python | 3.12+ (x86_64) | Core backend pipeline, ML inference, and API clients |
| **Web Presentation** | Streamlit | 1.40.0+ | Modern reactive web user interface & charts |
| **Bytecode Analysis** | Androguard | 4.1.4+ | DEX decompiler, manifest extraction, string pool parser |
| **Machine Learning** | Scikit-learn & XGBoost | 1.6+ / 2.1+ | Random Forest & Gradient Boosted Tree models |
| **Deep Learning** | PyTorch | 2.5.0+ | Multi-Branch TCN, BiLSTM, and Transformer inference |
| **Explainable AI** | SHAP | 0.46.0+ | Shapley additive feature importance computation |
| **Device Bridge** | Android ADB | 35.0.2 (Windows exe via ~/bin/adb wrapper) | Bridge connecting WSL2 scripts to Windows Android Emulator |
| **Mobile Virtualization** | Android Studio Emulator | API 34 ("Google APIs" system image, rooted) | Dynamic execution sandbox with strace & monkey |
| **Threat Intelligence** | AbuseIPDB REST API v2 | HTTPS JSON API | Live IP reputation & C2 confidence scoring |

---

## 9. Input & Output Data Specifications

### 9.1 System Inputs
1. **APK File:** Valid Android installation package (.apk) uploaded via browser or CLI.
2. **Analysis Duration:** Execution time slider for dynamic tracing ($10$ to $120$ seconds, default $30$s).
3. **API Credentials:** Optional AbuseIPDB API key provided via `.env` file, environment variable `ABUSEIPDB_API_KEY`, or dashboard input field.

### 9.2 System Outputs
1. **Static Analysis Verdict:**
   - Classification: `MALWARE` or `BENIGN`
   - Confidence percentage: e.g. `98.14% Malware Confidence`
   - Soft voting distribution: Individual Random Forest and XGBoost probabilities
2. **SHAP Local Explanation:**
   - Horizontal feature attribution chart ranking top evidence features
   - Feature taxonomy tags (Permission, Intent, API Call, Command)
3. **Threat Intelligence Audit:**
   - Table of embedded public IPs, associated country flags, and ISP names
   - Abuse Confidence Score gauge (0-100%) and total historical abuse reports
4. **Dynamic Trace Verdict:**
   - Syscall execution counts and frequency distribution
   - Top 5 dynamic syscall indicators (e.g. `ioctl`, `futex`, `epoll_wait`, `write`, `getuid32`)
5. **Structured JSON Audit Report (`apk_analysis_report.json`):** Machine-readable exportable schema for SIEM / SOC ingestion.

---

## 10. Limitations & Academic Integrity Disclosures

1. **Dataset Temporality:** The Drebin-215 dataset represents older Android architecture conventions. Modern Android APIs (Android 10+) enforce scoped storage and runtime permission grants that may not fully mirror older manifest-only permissions.
2. **Dynamic Evasion Resistance:** Highly sophisticated malware equipped with sandbox detection (e.g., checking `ro.kernel.qemu`, inspecting battery sensors, or timing loops) may stay dormant during the 30-second trace window.
3. **ABI Architecture Mapping:** CICMalDroid2020 training traces used 32-bit ARM CPU system call names. The live emulator running on modern x86_64 host hardware maps system calls via an internal alias translation table (`mmap` $\to$ `mmap2`, `lseek` $\to$ `_llseek`, `openat` $\to$ `open`).
"""

print("Generating Markdown documentation...")
with open(BASE_DIR / "PROJECT_TECHNICAL_SPECIFICATION.md", "w", encoding="utf-8") as f:
    f.write(MD_CONTENT)
print("Saved PROJECT_TECHNICAL_SPECIFICATION.md")


# -----------------------------------------------------------------------------
# 2. WORD DOCUMENT (DOCX) GENERATION
# -----------------------------------------------------------------------------
print("Generating Word Document (DOCX)...")
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = docx.Document()

# Page Margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Color Palette: Deep Navy (#1A365D), Slate Blue (#2B6CB0), Charcoal (#2D3748)
NAVY = RGBColor(26, 54, 93)
BLUE = RGBColor(43, 108, 176)
CHARCOAL = RGBColor(45, 55, 72)
MUTED = RGBColor(113, 128, 150)

# Configure default style
style_normal = doc.styles['Normal']
style_normal.font.name = 'Calibri'
style_normal.font.size = Pt(10.5)
style_normal.font.color.rgb = CHARCOAL

def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = NAVY
    p.paragraph_format.space_after = Pt(4)

def add_subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(13)
    run.font.italic = True
    run.font.color.rgb = BLUE
    p.paragraph_format.space_after = Pt(20)

def add_h1(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = NAVY
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)

def add_h2(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = BLUE
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)

def add_p(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    return p

def add_callout(text, title="NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.8)
    
    # Border & shading
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F0F4F8')
    tcPr.append(shd)
    
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(4)
    cp.paragraph_format.space_after = Pt(4)
    tr = cp.add_run(f"[{title}] ")
    tr.bold = True
    tr.font.color.rgb = BLUE
    cp.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_table(headers, rows):
    tbl = doc.add_table(rows=len(rows)+1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_cells[i], '1A365D')
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
    for r_idx, row_data in enumerate(rows):
        row_cells = tbl.rows[r_idx+1].cells
        bg_color = 'FFFFFF' if r_idx % 2 == 0 else 'F7FAFC'
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            if c_idx > 0 and any(char.isdigit() for char in str(val)):
                row_cells[c_idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

# Build Document Content
add_title("AI-Based Detection of Malicious Sideloaded Android Applications")
add_subtitle("Comprehensive System Architecture, Design Flow, Dataflow & Technical Specifications")

add_h1("Executive Summary")
add_p(
    "Sideloaded Android applications—installed outside official app stores like Google Play—bypass "
    "pre-deployment vetting, exposing mobile ecosystems to sophisticated adware, spyware, banking Trojans, "
    "and root exploits. Modern Android malware routinely employs obfuscation and dynamic payload loading "
    "that readily defeat signature-based static antiviruses."
)
add_p(
    "This system provides a holistic, multi-tier defense-in-depth architecture combining: (1) safe static APK "
    "decompilation and 215 Drebin feature extraction, (2) Random Forest and XGBoost soft voting ensemble "
    "achieving 98.14% test accuracy, (3) local Explainable AI via SHAP TreeExplainer, (4) automated runtime "
    "threat intelligence via AbuseIPDB API v2, (5) dynamic Android sandbox emulation capturing 139 Linux "
    "kernel system calls, and (6) a modern reactive Streamlit web dashboard."
)

add_callout(
    "Key Innovation: Dual static-dynamic synergy augmented by runtime threat intelligence and local Shapley "
    "feature attribution delivers both high detection accuracy (98.14% static / 97.31% dynamic) and verifiable "
    "human-in-the-loop auditability.",
    title="KEY HIGHLIGHT"
)

add_h1("1. System Architecture & High-Level Design Flow")
add_p(
    "The architecture is organized into five decoupled tiers ensuring fault tolerance, strict isolation, and "
    "high modularity:"
)
add_p(
    "1. Presentation Tier (Streamlit Web Dashboard): Provides reactive tabs for APK triage, model evaluation "
    "metrics, dynamic sandbox execution, and threat intelligence visualization.\n"
    "2. Static Ingestion & Bytecode Engine: Uses Androguard 4.x to inspect AndroidManifest.xml and classes.dex "
    "bytecode without executing untrusted code.\n"
    "3. Static ML & Explainability Engine: Evaluates the 215-dimensional Drebin binary vector across Random Forest "
    "and XGBoost, computing soft ensemble probabilities and SHAP TreeExplainer attributions.\n"
    "4. Threat Intelligence Engine: Harvests embedded public IPv4 addresses and domain hostnames, filters RFC 1918 "
    "private scopes, performs live DNS lookups, and queries AbuseIPDB v2.\n"
    "5. Dynamic Emulation Sandbox: Connects via ADB bridge to a rooted Android Virtual Device (AVD), injects the APK, "
    "fuzzes the UI with Monkey, and records 139 Linux kernel syscalls via strace."
)

add_h1("2. Dataflow Architecture (DFD Specifications)")
add_h2("2.1 DFD Level 0 (Context Level)")
add_p(
    "The external security analyst supplies an Android APK file and execution duration parameters. The system queries "
    "external AbuseIPDB threat intelligence feeds and outputs comprehensive verdicts, SHAP feature rankings, "
    "network risk scores, and downloadable JSON audit reports."
)

add_h2("2.2 DFD Level 1 (Subsystem Decomposition)")
add_p(
    "- Process 1.0 (Static Extraction): Parses APK ZIP containers, generating a 215-dimensional binary feature vector "
    "and extracting string pool indicators.\n"
    "- Process 2.0 (Threat Intel Enrichment): Filters public IPs and queries AbuseIPDB v2, returning abuse confidence "
    "scores and geographic metadata.\n"
    "- Process 3.0 (Static ML & SHAP): Feeds feature vectors to RF and XGBoost models, generating ensemble probabilities "
    "and Shapley attribution values.\n"
    "- Process 4.0 (Dynamic Emulation): Installs APK on emulator, runs UI Monkey, records strace output, normalizes "
    "syscall frequencies, and evaluates against dynamic XGBoost and TCN.\n"
    "- Process 5.0 (Reporting & UI Aggregation): Combines multi-modal outputs into reactive Streamlit components and "
    "exportable audit JSON."
)

add_h1("3. Machine Learning & Deep Learning Algorithms")
add_p(
    "The system leverages specialized models optimized for static tabular feature vectors and dynamic sequential "
    "system-call frequencies:"
)
add_h2("3.1 Static Analysis (Drebin-215)")
add_p(
    "- Random Forest: 100 de-correlated decision trees with bootstrap aggregation (bagging) and Gini impurity splits.\n"
    "- XGBoost: Gradient boosted decision trees using second-order Taylor expansion loss approximation, shrinkage (0.1), "
    "and depth-6 tree pruning.\n"
    "- Soft Voting Ensemble: Combines probabilities via P_ensemble = 0.5 * (P_RF + P_XGB), minimizing individual model "
    "variance and achieving 98.14% test accuracy."
)

add_h2("3.2 Dynamic Analysis (CICMalDroid2020: 11,520 Apps)")
add_p(
    "- XGBoost Classifier: Outperformed deep learning architectures on tabular syscall frequency counts (97.31% accuracy).\n"
    "- Multi-Branch Temporal Convolutional Network (TCN): Utilizes causal dilated 1D convolutions and residual blocks.\n"
    "- Bidirectional LSTM (BiLSTM): Recurrent network capturing forward and backward temporal dependencies.\n"
    "- Transformer: Multi-head self-attention encoder evaluating long-range syscall sequence interactions."
)

add_h1("4. Experimental Results & Performance Benchmarks")

add_h2("4.1 Static Analysis Performance (Drebin-215 Held-Out Test Set)")
static_headers = ["Model", "Accuracy", "Precision", "Recall", "F1-Score", "ROC-AUC"]
static_rows = [
    ["Random Forest", "97.11%", "96.02%", "91.55%", "0.9373", "0.9926"],
    ["XGBoost", "98.00%", "95.91%", "95.63%", "0.9577", "0.9961"],
    ["Fused Ensemble (RF + XGB)", "98.14%", "96.75%", "95.34%", "0.9604", "0.9952"]
]
create_table(static_headers, static_rows)

add_h2("4.2 Dynamic Analysis Performance (Random Split vs. Cluster-Held-Out)")
dyn_headers = ["Model Architecture", "Random Test Acc", "Random ROC-AUC", "Cluster-Held-Out Acc", "Cluster-Held-Out ROC-AUC"]
dyn_rows = [
    ["XGBoost", "97.31%", "0.9915", "91.99% ± 3.36%", "0.9600 ± 0.0223"],
    ["Transformer", "90.93%", "0.9691", "86.15% ± 6.62%", "0.9419 ± 0.0125"],
    ["TCN (Multi-Branch)", "94.53%", "0.9554", "84.99% ± 3.14%", "0.9013 ± 0.0409"],
    ["BiLSTM", "90.80%", "0.9523", "79.69% ± 9.25%", "0.8806 ± 0.0552"]
]
create_table(dyn_headers, dyn_rows)

add_callout(
    "Cluster-Held-Out cross validation holds out entire malware families to prevent family-level data leakage. "
    "XGBoost maintained superior generalization (91.99% accuracy) compared to sequence deep learning models.",
    title="VALIDATION RIGOR"
)

add_h2("4.3 McNemar Statistical Significance Tests")
mcnemar_headers = ["Comparison", "Only TCN Correct", "Only Other Correct", "p-value", "Significant (alpha 0.05/3)"]
mcnemar_rows = [
    ["TCN vs. BiLSTM", "1,481", "870", "0.0", "Yes"],
    ["TCN vs. Transformer", "758", "891", "0.0011", "Yes"],
    ["TCN vs. XGBoost", "587", "1,393", "0.0", "Yes"]
]
create_table(mcnemar_headers, mcnemar_rows)

add_h1("5. Embedded Benchmark Figures")
if (STATIC_OUT / "1_feature_importance.png").exists():
    add_h2("Static Feature Importance & Global Distributions")
    doc.add_picture(str(STATIC_OUT / "1_feature_importance.png"), width=Inches(6.2))
    p = doc.add_paragraph("Figure 1: Static Drebin-215 Feature Importance across Random Forest and XGBoost.")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)

if (STATIC_OUT / "2_confusion_matrices.png").exists():
    add_h2("Static Confusion Matrices")
    doc.add_picture(str(STATIC_OUT / "2_confusion_matrices.png"), width=Inches(6.2))
    p = doc.add_paragraph("Figure 2: Confusion Matrices for Static Drebin Classifiers.")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)

if (DYNAMIC_OUT / "3_roc_pr_curves.png").exists():
    add_h2("Dynamic ROC and Precision-Recall Curves")
    doc.add_picture(str(DYNAMIC_OUT / "3_roc_pr_curves.png"), width=Inches(6.2))
    p = doc.add_paragraph("Figure 3: Multi-Model ROC and Precision-Recall Curves on CICMalDroid2020.")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)

add_h1("6. Technical Specifications & Stack Details")
spec_headers = ["Component", "Technology / Tool", "Version", "Operational Role"]
spec_rows = [
    ["Host Environment", "Windows 11 / WSL 2 Ubuntu", "Ubuntu 24.04", "Virtualization, Python execution & dev"],
    ["Language & Runtime", "Python (x86_64)", "3.12+", "Core computational pipeline & API server"],
    ["Frontend UI", "Streamlit", "1.40+", "Reactive web dashboard & user console"],
    ["Static Decompiler", "Androguard", "4.1.4+", "APK manifest & DEX bytecode parsing"],
    ["ML Frameworks", "Scikit-Learn & XGBoost", "1.6+ / 2.1+", "Random Forest & Gradient Boosted trees"],
    ["Deep Learning", "PyTorch", "2.5+", "TCN, BiLSTM, and Transformer inference"],
    ["Explainability", "SHAP", "0.46+", "Shapley TreeExplainer feature attributions"],
    ["Device Bridge", "Android ADB", "35.0+", "WSL2 to Windows Android emulator bridge"],
    ["Dynamic Sandbox", "Android Studio Emulator", "API 34 (Google APIs)", "Rooted virtual device with strace & monkey"],
    ["Threat Intelligence", "AbuseIPDB REST API v2", "HTTPS JSON", "Live C2 IP reputation & risk scoring"]
]
create_table(spec_headers, spec_rows)

add_h1("7. Input & Output Technical Specifications")
add_h2("7.1 System Inputs")
add_p(
    "- File Input: Android Package (.apk) file (valid ZIP container with AndroidManifest.xml and classes.dex).\n"
    "- Dynamic Execution Duration: Integer slider value (10 to 120 seconds, default: 30s).\n"
    "- Threat Intel API Key: AbuseIPDB v2 authentication token (passed via environment variable or UI input)."
)

add_h2("7.2 System Outputs")
add_p(
    "- Security Verdict: Binary classification (BENIGN vs. MALWARE) accompanied by calibrated confidence percentage.\n"
    "- Soft Voting Probability Breakdown: Disaggregated Random Forest and XGBoost probabilities.\n"
    "- SHAP Feature Attribution Chart: Interactive horizontal bar graph highlighting top features pushing predictions toward malware or benign.\n"
    "- Threat Intelligence Report: Summary table of extracted public IPs, domain hostnames, ISP ownership, and Abuse Confidence Scores.\n"
    "- Dynamic System-Call Summary: Execution count table of 139 system calls with top dynamic indicator rankings.\n"
    "- Machine-Readable Audit JSON: Downloadable JSON export compliant with enterprise SIEM/SOC ingest formats."
)

docx_path = BASE_DIR / "Android_Malware_Detection_Technical_Specification.docx"
doc.save(str(docx_path))
print(f"Saved DOCX: {docx_path}")


# -----------------------------------------------------------------------------
# 3. PDF GENERATION VIA REPORTLAB
# -----------------------------------------------------------------------------
print("Generating PDF via ReportLab...")
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

pdf_path = BASE_DIR / "Android_Malware_Detection_Technical_Specification.pdf"

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_header_footer(num_pages)
            super().showPage()
        super().save()

    def draw_header_footer(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#718096"))
        
        # Suppress on cover page
        if self._pageNumber > 1:
            # Header
            self.drawString(54, letter[1] - 36, "Android Malware Detection System — Technical Specification")
            self.setStrokeColor(colors.HexColor("#E2E8F0"))
            self.setLineWidth(0.5)
            self.line(54, letter[1] - 42, letter[0] - 54, letter[1] - 42)
            
            # Footer
            self.line(54, 46, letter[0] - 54, 46)
            self.drawString(54, 32, "Confidential — Dual Static/Dynamic Android Security Pipeline")
            page_text = f"Page {self._pageNumber} of {page_count}"
            self.drawRightString(letter[0] - 54, 32, page_text)
            
        self.restoreState()

doc_pdf = SimpleDocTemplate(
    str(pdf_path),
    pagesize=letter,
    leftMargin=54,
    rightMargin=54,
    topMargin=54,
    bottomMargin=54
)

styles = getSampleStyleSheet()

# Custom styles
primary_color = colors.HexColor("#1A365D")   # Deep Navy
secondary_color = colors.HexColor("#2B6CB0") # Slate Blue
charcoal_color = colors.HexColor("#2D3748")  # Charcoal body text

title_style = ParagraphStyle(
    'DocTitle',
    parent=styles['Normal'],
    fontName='Helvetica-Bold',
    fontSize=22,
    leading=26,
    textColor=primary_color,
    alignment=1, # Center
    spaceAfter=8
)

subtitle_style = ParagraphStyle(
    'DocSubtitle',
    parent=styles['Normal'],
    fontName='Helvetica-Oblique',
    fontSize=12,
    leading=16,
    textColor=secondary_color,
    alignment=1,
    spaceAfter=20
)

h1_style = ParagraphStyle(
    'SectionH1',
    parent=styles['Heading1'],
    fontName='Helvetica-Bold',
    fontSize=14,
    leading=18,
    textColor=primary_color,
    spaceBefore=14,
    spaceAfter=6,
    keepWithNext=True
)

h2_style = ParagraphStyle(
    'SectionH2',
    parent=styles['Heading2'],
    fontName='Helvetica-Bold',
    fontSize=11,
    leading=15,
    textColor=secondary_color,
    spaceBefore=10,
    spaceAfter=4,
    keepWithNext=True
)

body_style = ParagraphStyle(
    'BodyTextCustom',
    parent=styles['Normal'],
    fontName='Helvetica',
    fontSize=9.5,
    leading=13.5,
    textColor=charcoal_color,
    spaceAfter=6
)

callout_style = ParagraphStyle(
    'CalloutText',
    parent=styles['Normal'],
    fontName='Helvetica',
    fontSize=9,
    leading=13,
    textColor=colors.HexColor("#1A202C")
)

table_text = ParagraphStyle(
    'TableText',
    parent=styles['Normal'],
    fontName='Helvetica',
    fontSize=8.5,
    leading=11,
    textColor=charcoal_color
)

table_header = ParagraphStyle(
    'TableHeader',
    parent=styles['Normal'],
    fontName='Helvetica-Bold',
    fontSize=8.5,
    leading=11,
    textColor=colors.white,
    alignment=1
)

story = []

# Title Section
story.append(Paragraph("AI-Based Detection of Malicious Sideloaded Android Applications", title_style))
story.append(Paragraph("System Architecture, Design Flow, Dataflow & Technical Specifications", subtitle_style))
story.append(HRFlowable(width="100%", thickness=1.5, color=secondary_color, spaceBefore=4, spaceAfter=14))

# Executive Summary
story.append(Paragraph("Executive Summary", h1_style))
story.append(Paragraph(
    "Sideloaded Android applications—installed outside official app stores like Google Play—bypass pre-deployment vetting, "
    "exposing users to sophisticated adware, spyware, banking Trojans, and root exploits. Modern Android malware routinely "
    "employs obfuscation and dynamic payload loading that readily defeat signature-based static antiviruses.",
    body_style
))
story.append(Paragraph(
    "This system provides a holistic, multi-tier defense-in-depth architecture combining: (1) safe static APK decompilation "
    "and 215 Drebin feature extraction, (2) Random Forest and XGBoost soft voting ensemble achieving <b>98.14% test accuracy</b>, "
    "(3) local Explainable AI via SHAP TreeExplainer, (4) automated runtime threat intelligence via AbuseIPDB API v2, "
    "(5) dynamic Android sandbox emulation capturing 139 Linux kernel system calls, and (6) a modern reactive Streamlit web dashboard.",
    body_style
))

# Callout box
callout_data = [[Paragraph("<b>KEY HIGHLIGHT:</b> Dual static-dynamic synergy augmented by runtime threat intelligence and local Shapley feature attribution delivers both high detection accuracy (98.14% static / 97.31% dynamic) and verifiable human-in-the-loop auditability.", callout_style)]]
callout_tbl = Table(callout_data, colWidths=[letter[0] - 108])
callout_tbl.setStyle(TableStyle([
    ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F0F4F8")),
    ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#CBD5E0")),
    ('PADDING', (0,0), (-1,-1), 8),
    ('BOTTOMPADDING', (0,0), (-1,-1), 8),
]))
story.append(callout_tbl)
story.append(Spacer(1, 10))

# 1. System Architecture
story.append(Paragraph("1. System Architecture & High-Level Design Flow", h1_style))
story.append(Paragraph(
    "The architecture is organized into five decoupled tiers ensuring fault tolerance, strict isolation, and high modularity:",
    body_style
))
arch_points = [
    "<b>1. Presentation Tier (Streamlit Web Dashboard):</b> Provides reactive tabs for single-click APK triage, model evaluation metrics, live dynamic sandbox execution, and threat intelligence visualization.",
    "<b>2. Static Ingestion & Bytecode Engine:</b> Uses Androguard 4.x to inspect AndroidManifest.xml and classes.dex bytecode without executing untrusted code.",
    "<b>3. Static ML & Explainability Engine:</b> Evaluates the 215-dimensional Drebin binary vector across Random Forest and XGBoost, computing soft ensemble probabilities and SHAP TreeExplainer attributions.",
    "<b>4. Threat Intelligence Engine:</b> Harvests embedded public IPv4 addresses and domain hostnames, filters RFC 1918 private scopes, performs live DNS lookups, and queries AbuseIPDB v2.",
    "<b>5. Dynamic Emulation Sandbox:</b> Connects via ADB bridge to a rooted Android Virtual Device (AVD), injects the APK, fuzzes the UI with Monkey, and records 139 Linux kernel syscalls via strace."
]
for pt in arch_points:
    story.append(Paragraph(f"• {pt}", body_style))

# 2. Machine Learning & Deep Learning Algorithms
story.append(Paragraph("2. Machine Learning & Deep Learning Algorithms", h1_style))
story.append(Paragraph("2.1 Static Analysis (Drebin-215)", h2_style))
story.append(Paragraph(
    "The static pipeline trains on 15,036 application vectors (5,560 malware, 9,476 benign) deduplicated to 7,254 unique samples across 215 binary indicators (permissions, intents, API calls, commands). "
    "An ensemble soft-voting scheme fuses Random Forest (100 trees, Gini impurity) and XGBoost (learning rate 0.1, depth 6) via arithmetic probability averaging: "
    "<i>P<sub>ensemble</sub> = 0.5 · (P<sub>RF</sub> + P<sub>XGB</sub>)</i>.",
    body_style
))

story.append(Paragraph("2.2 Dynamic Analysis (CICMalDroid2020)", h2_style))
story.append(Paragraph(
    "Dynamic analysis evaluates 11,520 applications traced during runtime execution across 139 Linux kernel system calls. "
    "Models include tabular XGBoost, Multi-Branch Temporal Convolutional Network (TCN with dilated causal convolutions), "
    "Bidirectional LSTM (BiLSTM), and Multi-Head Self-Attention Transformer. Crucially, evaluation utilizes "
    "<b>Cluster-Held-Out cross-validation</b>, holding out entire malware families to prevent cross-family data leakage.",
    body_style
))

# 3. Benchmark Tables
story.append(Paragraph("3. Experimental Results & Performance Benchmarks", h1_style))
story.append(Paragraph("Table 1: Static Analysis Evaluation (Drebin-215 Held-Out Test Split)", h2_style))

t1_headers = ["Model", "Accuracy", "Precision", "Recall", "F1-Score", "ROC-AUC"]
t1_rows = [
    ["Random Forest", "97.11%", "96.02%", "91.55%", "0.9373", "0.9926"],
    ["XGBoost", "98.00%", "95.91%", "95.63%", "0.9577", "0.9961"],
    ["Fused Ensemble (RF + XGB)", "98.14%", "96.75%", "95.34%", "0.9604", "0.9952"]
]
t1_table_data = [[Paragraph(h, table_header) for h in t1_headers]]
for row in t1_rows:
    t1_table_data.append([Paragraph(cell, table_text) for cell in row])

t1_tbl = Table(t1_table_data, colWidths=[140, 70, 70, 70, 70, 70])
t1_tbl.setStyle(TableStyle([
    ('BACKGROUND', (0,0), (-1,0), primary_color),
    ('ALIGN', (1,1), (-1,-1), 'CENTER'),
    ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E0")),
    ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F7FAFC")]),
    ('PADDING', (0,0), (-1,-1), 4),
]))
story.append(t1_tbl)
story.append(Spacer(1, 8))

story.append(Paragraph("Table 2: Dynamic System-Call Evaluation (CICMalDroid2020: Random vs. Cluster-Held-Out)", h2_style))
t2_headers = ["Model", "Random Acc", "Random ROC-AUC", "Cluster-Held-Out Acc", "Cluster-Held-Out ROC"]
t2_rows = [
    ["XGBoost", "97.31%", "0.9915", "91.99% ± 3.36%", "0.9600 ± 0.0223"],
    ["Transformer", "90.93%", "0.9691", "86.15% ± 6.62%", "0.9419 ± 0.0125"],
    ["TCN (Multi-Branch)", "94.53%", "0.9554", "84.99% ± 3.14%", "0.9013 ± 0.0409"],
    ["BiLSTM", "90.80%", "0.9523", "79.69% ± 9.25%", "0.8806 ± 0.0552"]
]
t2_table_data = [[Paragraph(h, table_header) for h in t2_headers]]
for row in t2_rows:
    t2_table_data.append([Paragraph(cell, table_text) for cell in row])

t2_tbl = Table(t2_table_data, colWidths=[130, 85, 85, 105, 85])
t2_tbl.setStyle(TableStyle([
    ('BACKGROUND', (0,0), (-1,0), primary_color),
    ('ALIGN', (1,1), (-1,-1), 'CENTER'),
    ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E0")),
    ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F7FAFC")]),
    ('PADDING', (0,0), (-1,-1), 4),
]))
story.append(t2_tbl)
story.append(Spacer(1, 8))

# 4. Embedded Visualizations
story.append(PageBreak())
story.append(Paragraph("4. Benchmark Figures & Empirical Evidence", h1_style))

if (STATIC_OUT / "1_feature_importance.png").exists():
    story.append(Paragraph("Figure 1: Static Feature Importance (Drebin-215)", h2_style))
    story.append(RLImage(str(STATIC_OUT / "1_feature_importance.png"), width=6.2*inch, height=2.8*inch))
    story.append(Spacer(1, 6))

if (DYNAMIC_OUT / "3_roc_pr_curves.png").exists():
    story.append(Paragraph("Figure 2: Dynamic ROC and Precision-Recall Curves (CICMalDroid2020)", h2_style))
    story.append(RLImage(str(DYNAMIC_OUT / "3_roc_pr_curves.png"), width=6.2*inch, height=2.6*inch))
    story.append(Spacer(1, 6))

if (STATIC_OUT / "2_confusion_matrices.png").exists():
    story.append(Paragraph("Figure 3: Confusion Matrices for Static Drebin Classifiers", h2_style))
    story.append(RLImage(str(STATIC_OUT / "2_confusion_matrices.png"), width=6.2*inch, height=2.5*inch))
    story.append(Spacer(1, 6))

# 5. Technical Specifications
story.append(PageBreak())
story.append(Paragraph("5. Technical Specifications & Stack Details", h1_style))

spec_headers = ["Layer", "Component", "Version", "Role"]
spec_rows = [
    ["Host Platform", "Windows 11 / WSL 2 Ubuntu", "24.04 LTS", "Virtualization, Python runtime & dev host"],
    ["Language", "Python (x86_64)", "3.12+", "Core computational pipeline & API clients"],
    ["Frontend UI", "Streamlit", "1.40+", "Reactive web dashboard & user console"],
    ["Decompiler", "Androguard", "4.1.4+", "APK manifest & DEX bytecode parsing"],
    ["Tabular ML", "Scikit-Learn & XGBoost", "1.6+ / 2.1+", "Random Forest & Gradient Boosted trees"],
    ["Deep Learning", "PyTorch", "2.5+", "TCN, BiLSTM, and Transformer inference"],
    ["Explainability", "SHAP", "0.46+", "Shapley TreeExplainer feature attributions"],
    ["Device Bridge", "Android ADB", "35.0+", "WSL2 to Windows Android emulator bridge"],
    ["Dynamic Sandbox", "Android Studio Emulator", "API 34 (Google APIs)", "Rooted virtual device with strace & monkey"],
    ["Threat Intel", "AbuseIPDB REST API v2", "HTTPS JSON", "Live C2 IP reputation & risk scoring"]
]
spec_table_data = [[Paragraph(h, table_header) for h in spec_headers]]
for row in spec_rows:
    spec_table_data.append([Paragraph(cell, table_text) for cell in row])

spec_tbl = Table(spec_table_data, colWidths=[90, 150, 80, 170])
spec_tbl.setStyle(TableStyle([
    ('BACKGROUND', (0,0), (-1,0), primary_color),
    ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E0")),
    ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F7FAFC")]),
    ('PADDING', (0,0), (-1,-1), 4),
]))
story.append(spec_tbl)
story.append(Spacer(1, 10))

# 6. Inputs and Outputs
story.append(Paragraph("6. Input & Output Specifications", h1_style))
story.append(Paragraph(
    "<b>Inputs:</b><br/>"
    "• Android Package (.apk) file: Valid ZIP archive containing AndroidManifest.xml and classes.dex bytecode.<br/>"
    "• Analysis Duration: Integer slider (10 to 120 seconds, default 30s).<br/>"
    "• Threat Intel API Key: Optional AbuseIPDB v2 key provided via .env, environment variable, or UI.",
    body_style
))
story.append(Spacer(1, 4))
story.append(Paragraph(
    "<b>Outputs:</b><br/>"
    "• Static Verdict Banner: High-visibility classification (BENIGN vs. MALWARE) with calibrated probability.<br/>"
    "• Soft Voting Distribution: Disaggregated probabilities from Random Forest and XGBoost.<br/>"
    "• SHAP Local Feature Attribution: Ranked horizontal bar chart identifying exact positive and negative drivers.<br/>"
    "• Threat Intelligence Report: Harvested public IPs, domains, Abuse Confidence Scores (0-100%), and report counts.<br/>"
    "• Dynamic Syscall Summary: Table of 139 system calls with top dynamic indicator rankings.<br/>"
    "• Audit Report (JSON): Downloadable <code>apk_analysis_report.json</code> for enterprise SIEM ingestion.",
    body_style
))

# 7. Limitations & Academic Integrity
story.append(Paragraph("7. Limitations & Academic Integrity Disclosures", h1_style))
story.append(Paragraph(
    "1. <b>Dataset Temporality:</b> Drebin-215 reflects historical Android permissions; modern Android runtime permissions may differ.<br/>"
    "2. <b>Anti-Analysis Evasion:</b> Advanced malware may detect emulator parameters (qemu, build properties) or sleep to evade dynamic trace.<br/>"
    "3. <b>Architecture Syscall Aliases:</b> CICMalDroid2020 used 32-bit ARM names. Live x86_64 emulator traces use internal alias mapping (mmap->mmap2, lseek->_llseek).",
    body_style
))

# Build PDF
doc_pdf.build(story, canvasmaker=NumberedCanvas)
print(f"Saved PDF: {pdf_path}")

print("\nAll 3 documentation formats generated successfully!")
